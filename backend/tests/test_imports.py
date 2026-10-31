"""客户导入：模板、解析、校验、跳过、commit。

本文件只钉“会静默写坏客户数据”的东西：重名必须跳过而不是覆盖、负责人必须
可解析、manager 不能跨部门、commit 必须单事务。
"""

from io import BytesIO

from openpyxl import load_workbook

from app.modules.imports.service import (
    CUSTOMER_TEMPLATE_COLUMNS,
    build_customer_template,
)


def test_template_has_headers_example_and_guide():
    workbook = load_workbook(BytesIO(build_customer_template()), read_only=True)

    assert workbook.sheetnames == ["客户", "填写说明"]
    rows = list(workbook["客户"].iter_rows(values_only=True))
    assert rows[0] == tuple(name for name, _ in CUSTOMER_TEMPLATE_COLUMNS)
    assert rows[0][0] == "客户名称*"
    assert rows[0][1] == "负责人*"
    assert rows[1][0] == "杭州示例科技有限公司"
    assert rows[1][1] == "sales1"
